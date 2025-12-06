from docx import Document

def create_sample_docx():
    doc = Document()
    doc.add_heading('Python 基础测试题', 0)

    # 题目 1
    doc.add_paragraph('1. Python 中用于输出内容的函数是？')
    doc.add_paragraph('A. input()')
    doc.add_paragraph('B. print()')
    doc.add_paragraph('C. output()')
    doc.add_paragraph('D. write()')
    doc.add_paragraph('答案：B')
    
    doc.add_paragraph('') # 空行

    # 题目 2
    doc.add_paragraph('2. 下列哪个不是 Python 的数据类型？')
    doc.add_paragraph('A. int')
    doc.add_paragraph('B. float')
    doc.add_paragraph('C. char')
    doc.add_paragraph('D. list')
    doc.add_paragraph('答案：C')

    doc.add_paragraph('') # 空行

    # 题目 3
    doc.add_paragraph('3. 如何定义一个函数？')
    doc.add_paragraph('A. func my_func():')
    doc.add_paragraph('B. define my_func():')
    doc.add_paragraph('C. def my_func():')
    doc.add_paragraph('D. function my_func():')
    doc.add_paragraph('答案：C')

    doc.save('sample_exam.docx')
    print("已生成示例试卷：sample_exam.docx")

if __name__ == "__main__":
    create_sample_docx()
